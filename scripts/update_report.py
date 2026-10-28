#!/usr/bin/env python3
"""Extend the user's edited DOCX, preserving its original sections and diagrams."""
import argparse
import csv
from collections import defaultdict
from pathlib import Path
from statistics import median

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parents[1]

def main():
    args=argparse.ArgumentParser()
    args.add_argument('--source',type=Path,required=True)
    args.add_argument('--output',type=Path,required=True)
    args.add_argument('--results',type=Path,default=ROOT/'results/formal')
    args.add_argument('--github-url',default='')
    a=args.parse_args()
    if a.source.resolve()==a.output.resolve(): raise ValueError('Preserve the original document: choose a new output')
    d=Document(a.source)
    original=list(d.paragraphs)
    # These replacements are anchored to exact paragraph starts, not page numbers.
    replacements={
      '本次未自行修改算法源码或课程脚本。':
        '基础实验未自行修改算法源码或课程脚本。GateANN 使用课程 AIO 修补，将 cb.data 改为 cb[j].data，以正确绑定异步请求回调，不改变算法或索引布局；本地依赖通过解压软件包提供。新增优化阶段采用独立 C++ 驱动复用原 GateANN 库，不改写基线搜索逻辑。基础正式实验标识为 20261010-190118；优化对照为单独保存的 2026-10-28 新一批运行，见第 5.9—5.12 节。',
      '针对 GateANN Low 档候选不足的现象，首先通过补充实验':
        '针对 GateANN Low 档候选不足的现象，先通过改变 L 的补充实验识别质量与穿越成本的关系，再实现“标签倒排枚举—PQ 预选—完整向量精确重排”自适应路径。第 5.1—5.4 节保留前期证据，第 5.5—5.8 节说明方案与验证标准，第 5.9—5.12 节给出已完成的代码、重新测量的三轮对照、调试记录和代码提交位置。',
      '拟增加一条自适应候选生成路径。该方案尚未实现':
        '已增加一条自适应候选生成路径，具体实现位于 src/adaptive_search.h，统一实验入口为 src/benchmark.cpp。实现边界为课程静态 uint8/L2 索引、单标签谓词、单查询线程和 identity ID/tag/page 映射；不支持的索引布局或查询格式显式报错。以下五项设计均已落实并验证。',
      '（1）建立标签倒排列表与计数。':
        '（1）从完整 Base 的 CSR 标签文件建立全部标签的倒排列表与计数，以 uint64 offsets 和 uint32 向量 ID 扁平存储于 DRAM。忽略值为零的项，同一向量的重复标签去重。列表容量和构建时间单独统计，不使用查询 Ground Truth 构建候选。',
      '（2）按选择率自适应路由。':
        '（2）按选择率自适应路由。每次查询以倒排表计数除以 Base 总量计算选择率；严格低于阈值时走枚举路径，其余调用原 GateANN pipe_search。代表阈值为 1%；增加阈值为 0、0.5% 的对照，检验路由作用。',
      '（3）使用 PQ 预选候选。':
        '（3）复用原 GateANN 索引内存 PQ 数据与查询距离表，分批计算满足标签的全部候选的 PQ 近似距离，以 partial_sort 保留前 M 个。每批最多 MAX_N_EDGES=1024，避免越过原库 scratch 上限；不足 M 时保留全部候选。实际扫描 M=10/20/40/80/160，代表配置选 M=40。',
      '（4）读取完整向量并精确重排。':
        '（4）每批最多 8 个 Linux AIO 请求，从原 SSD 图记录读取入选候选的完整向量，按精确 uint8 L2 距离重排并返回 top-10。距离相等时以 ID 确定顺序。不把 Base 完整向量常驻 DRAM；同页重复请求仍按请求数计入逻辑 I/O。无候选时输出无效 ID 哨兵，候选不足 K 时其余位置保持哨兵。',
      '（5）分阶段记录查询成本。':
        '（5）逐查询记录路由、PQ 扫描/预选、AIO 及精确重排耗时，并记录枚举候选数、重排数、逻辑 I/O 和路由数量。倒排表在启动阶段由标签文件重建并常驻 DRAM，不新增持久化文件；初始化构建时间、内存容量及进程峰值 RSS 分开报告。',
      '预期收益主要针对 Low 档：':
        '设计目标是改善 Low 档的有效候选覆盖，并降低接近 Recall 下的查询成本。实测在接近 70% Recall 区间，M=10 达到 72.75% Recall，延迟明显低于原 L=1280；M=40 将 Recall 提升至 98.35%。Medium/High 在阈值 1% 下全部回退原图搜索。具体三轮结果与范围见第 5.10 节，不能将预实验单轮数据当作独立泛化结论。',
      '对照组为未改动的 GateANN L=320、640、1280；实验组':
        '验证分为前期预算补充实验和新增实现对照。前期对照为 GateANN L=320/640/1280；实现阶段另测原图搜索 L=320/1280 与新增自适应路径。统一 Formal 数据、原 SSD/PQ 索引、AIO、K=10、线程=1、beamwidth=8、FullAdj=32、mem_L=0、cache_budget=0。每配置独立运行 3 次，基线与优化使用同一驱动的直接查询计时，不与 10 月 10 日旧实验混合统计。',
      '参数验证计划先固定阈值考察 M，':
        '已先以单轮预实验检查实现和参数范围，再固定阈值 1% 考察 M=10/20/40/80/160；固定 M=40 考察阈值 0、0.5%、1%。另外分别对 Medium/High 测原图搜索与自适应路径。共 13 组配置、每组 3 轮、合计 39 次新进程运行。每轮交替反转配置顺序，减少固定顺序偏差。预实验仅用于选择参数，不计入三轮汇总。',
      '验证时还应记录附加索引空间及构建时间。':
        '本次已记录新增内存容量、启动构建时间和逐标签指标，并按上段原定目标检查。M=40 达到 Low 的 Recall 与延迟目标；Medium/High 的 Recall 未下降，延迟中位数增幅均小于 10%。这些阈值是设计验收线，不是统计显著性检验。参数选择和评估使用同一课程工作负载，后续需在独立标签/查询集上验证泛化。',
      '在本组数据中，DRAM 相对充足且过滤较弱时，':
        '在本组数据中，GateANN 原图搜索在过滤较弱时具有较低查询延迟；强过滤时，候选生成策略更影响 Recall。本文已实现基于标签倒排表的自适应 PQ 预选与 SSD 精确重排。在新批次 Low 对照中，原 GateANN L=1280 的 Recall 为 70.05%、延迟为 3.202 ms；新增 M=10 为 72.75%、0.553 ms，接近 Recall 时延迟降低约 82.7%。M=40 可达 98.35%、1.083 ms，但需要更多逻辑读取与额外倒排内存。',
    }
    found=set()
    for p in original:
        for start,new in replacements.items():
            if p.text.startswith(start):
                p.text=new; found.add(start); break
    if len(found)!=len(replacements): raise ValueError(f'Source text changed: missing {set(replacements)-found}')
    # Fix accidental doubled shell continuation characters in the edited source.
    for p in d.paragraphs:
        if p.text.endswith('\\\\'): p.text=p.text[:-1]
    anchor=next(p for p in d.paragraphs if p.text=='6 实验总结')
    def paragraph(text='',style=None): return anchor.insert_paragraph_before(text,style)
    def heading(text): return paragraph(text,'Heading 2')
    def table(caption,headers,rows):
        paragraph(caption)
        t=d.add_table(rows=1,cols=len(headers)); t.style='Table Grid'; t.autofit=True
        for c,text in zip(t.rows[0].cells,headers): c.text=str(text)
        for values in rows:
            for c,text in zip(t.add_row().cells,values): c.text=str(text)
        repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
        anchor._p.addprevious(t._tbl)
        paragraph()
    summaries=list(csv.DictReader((a.results/'summary.csv').open()))
    def get(bucket,policy,L=320,M=40,T=.01):
        return next(r for r in summaries if r['bucket']==bucket and r['policy']==policy and
                    int(r['L'])==L and int(r['M'])==M and float(r['threshold'])==T)
    def value(r,field,factor=1,decimals=3): return f"{float(r[field+'_median'])*factor:.{decimals}f}"
    def interval(r,field,factor=1,decimals=3):
        return value(r,field,factor,decimals)+f"\n[{float(r[field+'_min'])*factor:.{decimals}f}, {float(r[field+'_max'])*factor:.{decimals}f}]"

    heading('5.9 优化代码实现与运行边界')
    paragraph('实现采用增量扩展：原 GateANN 静态库提供索引加载、内存 PQ、图搜索与 AIO；新增代码实现 LabelPostings、选择率路由及枚举重排。这样可让基线和优化在同一驱动下测量，同时保留原系统作为回退路径。未重新训练 PQ、未重建 SSD 基础图、未把原始答案或 Ground Truth 用于搜索。Ground Truth 仅在整批 Query Phase 结束后打开用于评估，评估时还检查输出 ID 是否重复、越界或不满足标签。')
    paragraph('核心伪代码：count ← postings[label].size；若 count/N < threshold，则对 postings[label] 分批计算原 PQ 近似距离、取前 min(M,count) 个，按每批 8 个 AIO 请求读取原磁盘图记录并以完整向量精确 L2 重排；否则调用原 pipe_search。实际 C++ 实现、统一驱动和测试均位于待上传代码目录，不以伪代码替代实现。')
    paragraph('静态倒排构建复杂度约为 O(N+标签数+有效标签项数)，常驻空间 O(标签数+有效标签项数)；一条枚举查询扫描 C 个匹配候选，其 PQ 计算随 C 和 PQ 分块数增长，部分排序约 O(C log M)，精确重排读取 min(M,C) 条图记录。M 过大时 I/O 与精确重排会主导延迟。当前只验证课程单标签、uint8/L2、单线程和 identity 映射；不覆盖在线更新、多标签交集及重排过的索引。')
    heading('5.10 新增三轮对照结果与分析')
    paragraph('新增实验于 2026-10-28 在同一服务器独立运行。开始时负载均值约 2.01/2.82/3.26，和原始基础实验时期不同。原图搜索基线及优化均由同一 C++ 驱动调用，每条查询直接计时，整批耗时计算 QPS；初始化、倒排构建及读取 GT 的开销不计入 Query Phase。表中为三轮中位数，方括号为最小值和最大值；所有查询共用已加载索引，未主动清理 OS 缓存，程序向 SSD 图记录发出 Linux AIO 对齐读取。本批直接计时与基础实验按 QPS 换算的口径分别保存。')
    rows=[]
    for L in (320,1280):
        r=get('low','graph',L=L)
        rows.append([f'原图 L={L}',value(r,'recall_percent',decimals=2),interval(r,'mean_latency_us',.001),value(r,'qps',decimals=1),value(r,'mean_ios',decimals=3)])
    for M in (10,20,40,80,160):
        r=get('low','adaptive',M=M)
        rows.append([f'自适应 M={M}',value(r,'recall_percent',decimals=2),interval(r,'mean_latency_us',.001),value(r,'qps',decimals=1),value(r,'mean_ios',decimals=3)])
    table('表 9  Low 新批次搜索质量与查询成本（阈值 1%）',
          ['配置','Recall@10 %','平均延迟 ms\n中位数及范围','QPS','逻辑 I/O'],rows)
    p=paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(a.results/'optimization_comparison.png'),width=Inches(5.65))
    paragraph('图 5  Low 优化路径的 Recall–Latency 与阶段耗时；误差线为三轮延迟范围。')
    b=get('low','graph',L=1280); main=get('low','adaptive'); near=get('low','adaptive',M=10)
    gain=(1-float(near['mean_latency_us_median'])/float(b['mean_latency_us_median']))*100
    main_gain=(1-float(main['mean_latency_us_median'])/float(b['mean_latency_us_median']))*100
    paragraph(f'相近质量测点中，原 L=1280 为 70.05% Recall、3.202 ms，自适应 M=10 为 72.75%、0.553 ms，延迟降低约 {gain:.1f}%。这不是严格相同 Recall 下的插值比较，但优化同时取得更高质量与更低延迟。代表配置 M=40 达到 98.35%、1.083 ms，相对同批 L=1280 延迟降低约 {main_gain:.1f}%，Recall 高 28.30 个百分点；相对原 L=320 的 34.30% Recall，质量明显改善而延迟接近。因此现有优化已经形成清晰收益，无需把方案改成单纯增大搜索预算。')
    paragraph('机制上，M=40 每条查询平均枚举 5,490 个已匹配候选，并精确重排 40 个；不再沿不匹配节点做 Graph Tunneling。PQ 扫描/预选中位耗时约 '+value(main,'pq_us')+' μs，AIO 与精确重排约 '+value(main,'rerank_us')+' μs。注意这不是减少所有 I/O：原 L=1280 只有约 13.945 次逻辑 I/O，M=40 为 40 次；优化以额外有效候选读取换取更高覆盖，并避免大量穿越。M 增至 160 虽达 100% Recall，延迟升至 3.698 ms，说明重排预算仍存在质量—成本权衡。')
    rows=[]
    for bucket,label in [('medium','Medium'),('high','High')]:
        for policy,name in [('graph','原图'),('adaptive','自适应')]:
            r=get(bucket,policy)
            rows.append([label+' '+name,value(r,'recall_percent',decimals=2),interval(r,'mean_latency_us',.001),int(float(r['enumerated_median']))])
    table('表 10  Medium/High 回退原图搜索的结果（L=320，M=40，阈值 1%）',
          ['配置','Recall@10 %','平均延迟 ms\n中位数及范围','枚举查询数/200'],rows)
    changes={bucket:(float(get(bucket,'adaptive')['mean_latency_us_median'])/float(get(bucket,'graph')['mean_latency_us_median'])-1)*100 for bucket in ('medium','high')}
    paragraph(f'Medium/High 在 1% 阈值下 200 条查询均回退原图搜索，Recall 分别保持 72.05%/93.25%；延迟中位数增幅约 {changes["medium"]:.2f}%/{changes["high"]:.2f}%，在原定 10% 验收线上限内。逐轮返回 ID 与对应基线逐字节相同，验证回退逻辑没有改变检索结果；小幅时间差也可能包含缓存和后台噪声，不能仅据此归因于路由开销。')
    rows=[]
    for t in (0,.005,.01):
        r=get('low','adaptive',T=t)
        rows.append([f'{t*100:g}%',int(float(r['enumerated_median'])),value(r,'recall_percent',decimals=2),interval(r,'mean_latency_us',.001)])
    table('表 11  固定 M=40 的 Low 路由阈值对照',
          ['阈值','枚举查询数/200','Recall@10 %','平均延迟 ms\n中位数及范围'],rows)
    paragraph('阈值 0 完全关闭枚举，其结果 ID 与原图搜索相同；0.5% 仅路由 80/200 条查询，Recall 为 64.40%；1% 路由全部 Low 查询，Recall 为 98.35%。这组单变量对照支持候选路径切换确实改变了结果覆盖。阈值 1% 来自同一课程负载的试验选择，不能声称适合任意标签分布。')
    perlabel=defaultdict(list)
    for r in csv.DictReader((a.results/'per_label.csv').open()): perlabel[int(r['label'])].append(r)
    rows=[]
    for label,rs in sorted(perlabel.items()):
        rows.append([label,int(float(rs[0]['candidates'])),f'{median(float(r["recall_percent"]) for r in rs):.2f}',f'{median(float(r["query_us"]) for r in rs)/1000:.3f}'])
    table('表 12  代表配置 M=40 的 Low 逐标签结果（各 40 条查询，三轮中位数）',
          ['标签 ID','候选数/查询','Recall@10 %','平均延迟 ms'],rows)
    paragraph('五个 Low 标签均达到 98.25%—98.50% Recall，改善并非只来自某一个标签。匹配数量从 2,499 到 8,513，延迟随候选扫描规模总体增长；但标签语义、向量分布和测量波动仍会共同影响时间，不能把逐标签差异全部解释为候选数的因果效应。')
    paragraph('新增倒排表分配容量为 44,934,300 字节，即 42.85 MiB；M=40/Low 的启动构建时间为 '+interval(main,'postings_build_ms',decimals=2).replace('\n',' ')+' ms。对应整个进程峰值 RSS 为 '+interval(main,'peak_rss_kb',1/1024,2).replace('\n',' ')+' MiB，原 L=320 为 '+interval(get('low','graph'),'peak_rss_kb',1/1024,2).replace('\n',' ')+' MiB。仅统计容器容量不能代替整个进程峰值；构建临时缓冲、原库内存及分配器均可能影响 RSS。本实现每次启动重建，不产生新增持久化索引，不能据此宣称新增空间为零。以本批 Low M=40 相对 L=1280 的中位查询时间差估算，单次倒排构建成本需要约 131 次查询摊销；该估算未代替实际端到端服务测试，仍应按服务生命周期评估。')
    heading('5.11 调试记录与正确性验证')
    table('表 13  本次实现的调试与核查记录',['问题或检查','处理过程及结果'],[
        ['本地编译找不到 cblas.h','沿用基础实验解压的 OpenBLAS/Eigen；设置 CPLUS_INCLUDE_PATH。构建脚本读取官方 target 的原始 ABI 宏及链接参数，编译成功。'],
        ['元数据与路由边界','原生 C++ 测试覆盖 CSR 倒排构建、重复/零值标签、未知标签、阈值严格小于比较、零阈值关闭路由及截断输入；全部通过。'],
        ['精确重排是否读取正确向量','Debug Low 前 5 条独立查询，M 足够大以重排全部匹配候选，Recall@10=100%；输出还通过标签、ID 范围与去重检查。'],
        ['回退逻辑是否改变答案','Medium/High 三轮及 Low 阈值为 0 的三轮，共 9 组配对结果 ID 逐字节一致。'],
        ['优化是否只减少了测量工作','Ground Truth 在搜索结束后才读取；索引加载与倒排构建另计。原图基线同批重测，Recall 与课程基线一致。39 次运行全部正常结束，原始命令与资源日志保留。']])
    paragraph('实现仍有明确局限：仅以课程静态数据测试，没有在线增删标签、并发吞吐或独立标签测试集；同页候选未合并 AIO 读取，较大 M 下精确重排已成为主要成本。若需要进一步优化，可先对重排候选按扇区去重并批量读取，再在同等 Recall 下验证收益和新增分组开销；现有明显改善已经由实际代码与对照数据支持，不把这一后续方向写成已实现功能。')
    heading('5.12 优化实现代码（GitHub）与复现入口')
    github=a.github_url.strip()
    paragraph('GitHub 仓库：'+(github if github else '【待填写真实仓库链接：代码尚未发布，上传后替换此处】'))
    paragraph('代码包对应上述 GitHub 仓库，本地整理目录为 github_submission/。src/ 包含优化实现与统一驱动，tests/ 包含边界及 Debug 正确性测试，scripts/ 包含构建、三轮对照、统计绘图与本报告更新脚本；baseline/ 保留课程原始实验代码及来源，results/formal/ 保留 39 次新运行的 command.json、run.log、resources.txt、queries.csv、结果 ID、汇总表和输入/源码哈希。数据集、第三方源码与二进制不提交。')
    paragraph('代码包根目录的主要复现命令如下，COURSE_ROOT 指向已准备好的课程仓库：')
    for command in ['python3 scripts/build.py --gateann "$COURSE_ROOT/systems/GateANN-public"',
                    'python3 scripts/build.py --gateann "$COURSE_ROOT/systems/GateANN-public" --tests',
                    'python3 tests/check_correctness.py --course-root "$COURSE_ROOT"',
                    'python3 scripts/run_experiments.py --course-root "$COURSE_ROOT" --out results/reproduced --repeats 3',
                    'python3 scripts/analyze.py --results results/reproduced']:
        paragraph(command)

    # Replace stale cached page numbers by a native Word TOC field. Cached entries
    # deliberately have no old page numbers; Word can regenerate actual pagination.
    toc_title=next(p for p in d.paragraphs if p.text.strip()=='目录')
    # In the user's source, TOC occupies original paragraphs 5..32 (end field included).
    for p in original[5:33]:
        parent=p._p.getparent()
        if parent is not None: parent.remove(p._p)
    toc_start=toc_title.insert_paragraph_before('')
    toc_title._p.addnext(toc_start._p)
    begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin'); begin.set(qn('w:dirty'),'true')
    instruction=OxmlElement('w:instrText'); instruction.set(qn('xml:space'),'preserve'); instruction.text=' TOC \\o "1-2" \\h \\z \\u '
    separate=OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'),'separate')
    for element in (begin,instruction,separate): toc_start.add_run()._r.append(element)
    last=toc_start
    for number,h in enumerate([p for p in d.paragraphs if p.style.name in ('Heading 1','Heading 2')],1):
        bookmark=f'UpdatedHeading{number}'
        start=OxmlElement('w:bookmarkStart'); start.set(qn('w:id'),str(5000+number)); start.set(qn('w:name'),bookmark)
        end=OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'),str(5000+number))
        h._p.insert(1,start); h._p.append(end)
        toc_style_id='TOC1' if h.style.name=='Heading 1' else 'TOC2'
        toc_style=next((s for s in d.styles if s.style_id==toc_style_id),d.styles['Normal'])
        p=d.add_paragraph(style=toc_style)
        link=OxmlElement('w:hyperlink'); link.set(qn('w:anchor'),bookmark)
        r=OxmlElement('w:r'); text=OxmlElement('w:t'); text.text=h.text; r.append(text); link.append(r)
        p._p.append(link); last._p.addnext(p._p); last=p
    end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end'); last.add_run()._r.append(end)
    settings=d.settings.element
    update=settings.find(qn('w:updateFields'))
    if update is None: update=OxmlElement('w:updateFields'); settings.append(update)
    update.set(qn('w:val'),'true')

    # Explicit body formatting: 五号=10.5 pt, 宋体, 1.25 multiples.
    # Chinese 小四 is 12 pt; do not confuse it with 五号.
    for style in d.styles:
        if hasattr(style,'font'):
            style.font.name='宋体'; style.font.size=Pt(10.5)
            rpr=style.element.get_or_add_rPr(); fonts=rpr.find(qn('w:rFonts'))
            if fonts is not None:
                for key in list(fonts.attrib):
                    if key.endswith('Theme') or key.endswith('theme'): del fonts.attrib[key]
                fonts.set(qn('w:eastAsia'),'宋体')
        if hasattr(style,'paragraph_format'):
            style.paragraph_format.line_spacing=1.25
    roots=[d.element]
    for section in d.sections:
        roots.extend([section.header._element,section.footer._element])
    for root in roots:
        for p in root.iter(qn('w:p')):
            pp=p.get_or_add_pPr()
            spacing=pp.find(qn('w:spacing'))
            if spacing is None: spacing=OxmlElement('w:spacing'); pp.append(spacing)
            spacing.set(qn('w:line'),'300'); spacing.set(qn('w:lineRule'),'auto')
            grid=pp.find(qn('w:snapToGrid'))
            if grid is None: grid=OxmlElement('w:snapToGrid'); pp.append(grid)
            grid.set(qn('w:val'),'false')
        for r in root.iter(qn('w:r')):
            rp=r.get_or_add_rPr(); fonts=rp.find(qn('w:rFonts'))
            if fonts is None: fonts=OxmlElement('w:rFonts'); rp.insert(0,fonts)
            for key in list(fonts.attrib):
                if key.endswith('Theme') or key.endswith('theme'): del fonts.attrib[key]
            for key in ('ascii','hAnsi','eastAsia','cs'): fonts.set(qn('w:'+key),'宋体')
            for name in ('sz','szCs'):
                size=rp.find(qn('w:'+name))
                if size is None: size=OxmlElement('w:'+name); rp.append(size)
                size.set(qn('w:val'),'21')
    # Remove fixed row heights that could clip the now standardized body text.
    for table_obj in d.tables:
        for row in table_obj.rows:
            props=row._tr.get_or_add_trPr()
            for height in list(props.findall(qn('w:trHeight'))): props.remove(height)
    d.core_properties.title='基于 SSD 的过滤向量检索系统性能测试与设计比较（含优化实现）'
    a.output.parent.mkdir(parents=True,exist_ok=True)
    d.save(a.output)
    print(a.output)

if __name__=='__main__': main()
